"""Gherkin step implementations for slide comments feature."""

from __future__ import annotations

import os

from behave import given, then, when
from helpers import saved_pptx_path, test_pptx

from pptx import Presentation


# given ===================================================


@given("a slide with no comments")
def given_a_slide_with_no_comments(context):
    context.prs = Presentation(test_pptx("sld-comments"))
    context.slide = context.prs.slides[1]


@given("a slide with comments")
def given_a_slide_with_comments(context):
    context.prs = Presentation(test_pptx("sld-comments"))
    context.slide = context.prs.slides[0]


# when ====================================================


@when('I add a comment "{text}" by "{author}"')
def when_I_add_a_comment(context, text, author):
    context.slide.add_comment(text, author)


@when("I save and reload the presentation and get the first slide")
def when_I_save_and_reload_the_presentation_and_get_first_slide(context):
    context.prs.save(saved_pptx_path)
    context.prs = Presentation(saved_pptx_path)
    context.slide = context.prs.slides[0]
    os.remove(saved_pptx_path)


# then ====================================================


@then("slide.comments is an empty list")
def then_slide_comments_is_an_empty_list(context):
    comments = context.slide.comments
    assert comments == [], "expected empty list, got %s" % comments


@then("slide.comments has {count:d} items")
def then_slide_comments_has_count_items(context, count):
    comments = context.slide.comments
    assert len(comments) == count, "expected %d comments, got %d" % (count, len(comments))


@then('the first comment text is "{text}"')
def then_the_first_comment_text_is(context, text):
    actual = context.slide.comments[0].text
    assert actual == text, "expected '%s', got '%s'" % (text, actual)


@then('the first comment author is "{author}"')
def then_the_first_comment_author_is(context, author):
    actual = context.slide.comments[0].author
    assert actual == author, "expected '%s', got '%s'" % (author, actual)


@then("the first comment has a timestamp")
def then_the_first_comment_has_a_timestamp(context):
    ts = context.slide.comments[0].timestamp
    assert ts is not None, "expected a timestamp, got None"


@then('the second comment text is "{text}"')
def then_the_second_comment_text_is(context, text):
    actual = context.slide.comments[1].text
    assert actual == text, "expected '%s', got '%s'" % (text, actual)


@then('the second comment author is "{author}"')
def then_the_second_comment_author_is(context, author):
    actual = context.slide.comments[1].author
    assert actual == author, "expected '%s', got '%s'" % (author, actual)

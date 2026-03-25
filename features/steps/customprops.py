"""Gherkin step implementations for custom properties-related features."""

from __future__ import annotations

from behave import given, then, when
from behave.runner import Context
from helpers import saved_pptx_path

from pptx import Presentation


# given ===================================================


@given("I have a new presentation")
def step_given_new_presentation(context: Context):
    context.prs = Presentation()


# when ====================================================


@when('I set a custom property "{name}" to string "{value}"')
def step_when_set_custom_prop_string(context: Context, name: str, value: str):
    context.prs.custom_properties[name] = value


@when('I set a custom property "{name}" to int {value:d}')
def step_when_set_custom_prop_int(context: Context, name: str, value: int):
    context.prs.custom_properties[name] = value


@when('I set a custom property "{name}" to bool {value}')
def step_when_set_custom_prop_bool(context: Context, name: str, value: str):
    context.prs.custom_properties[name] = value.lower() == "true"


@when('I delete the custom property "{name}"')
def step_when_delete_custom_prop(context: Context, name: str):
    del context.prs.custom_properties[name]


# then ====================================================


@then('the custom property "{name}" of the saved presentation is "{value}"')
def step_then_custom_prop_string_value(context: Context, name: str, value: str):
    prs = Presentation(saved_pptx_path)
    assert prs.custom_properties[name] == value


@then('the custom property "{name}" of the saved presentation is int {value:d}')
def step_then_custom_prop_int_value(context: Context, name: str, value: int):
    prs = Presentation(saved_pptx_path)
    assert prs.custom_properties[name] == value


@then('the custom property "{name}" of the saved presentation is bool {value}')
def step_then_custom_prop_bool_value(context: Context, name: str, value: str):
    prs = Presentation(saved_pptx_path)
    expected = value.lower() == "true"
    assert prs.custom_properties[name] is expected


@then('the custom property "{name}" does not exist in the saved presentation')
def step_then_custom_prop_does_not_exist(context: Context, name: str):
    prs = Presentation(saved_pptx_path)
    assert name not in prs.custom_properties

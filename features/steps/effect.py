"""Gherkin step implementations for ShadowFormat-related features."""

from __future__ import annotations

from behave import given, then, when
from helpers import test_pptx

from pptx import Presentation
from pptx.util import Emu

# given ====================================================


@given("a ShadowFormat object that {inherits} as shadow")
def given_a_ShadowFormat_object_that_inherits_or_not(context, inherits):
    shape_idx = {"inherits": 0, "does not inherit": 1}[inherits]
    shape = Presentation(test_pptx("dml-effect")).slides[0].shapes[shape_idx]
    context.shadow = shape.shadow


@given("a shape with an outer shadow as shadow_shape")
def given_a_shape_with_an_outer_shadow(context):
    # Shape at index 2 has an explicit outer shadow with known values:
    # blurRad=40000, dist=23000, dir=5400000 (90.0 degrees), color=000000
    shape = Presentation(test_pptx("dml-effect")).slides[0].shapes[2]
    context.shadow = shape.shadow


# when =====================================================


@when("I assign {value} to shadow.inherit")
def when_I_assign_value_to_shadow_inherit(context, value):
    context.shadow.inherit = eval(value)


# then =====================================================


@then("shadow.inherit is {bool_str}")
def then_shadow_inherit_is_bool_val(context, bool_str):
    expected_value = eval(bool_str)
    actual_value = context.shadow.inherit
    assert actual_value is expected_value, "shadow.inherit is %s" % actual_value


@then('shadow.shadow_type is "{expected_type}"')
def then_shadow_type_is(context, expected_type):
    actual = context.shadow.shadow_type
    assert actual == expected_type, "shadow.shadow_type is %s" % actual


@then("shadow.angle is {expected_angle}")
def then_shadow_angle_is(context, expected_angle):
    expected = float(expected_angle)
    actual = context.shadow.angle
    assert actual == expected, "shadow.angle is %s" % actual


@then("shadow.blur_radius is {expected_emu}")
def then_shadow_blur_radius_is(context, expected_emu):
    expected = Emu(int(expected_emu))
    actual = context.shadow.blur_radius
    assert actual == expected, "shadow.blur_radius is %s" % actual


@then("shadow.distance is {expected_emu}")
def then_shadow_distance_is(context, expected_emu):
    expected = Emu(int(expected_emu))
    actual = context.shadow.distance
    assert actual == expected, "shadow.distance is %s" % actual


@then('shadow.color.rgb is "{expected_rgb}"')
def then_shadow_color_rgb_is(context, expected_rgb):
    actual = str(context.shadow.color.rgb)
    assert actual == expected_rgb, "shadow.color.rgb is %s" % actual

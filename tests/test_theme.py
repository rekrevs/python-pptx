"""Unit-test suite for `pptx.theme` module."""

from __future__ import annotations

from io import BytesIO

import pytest

from pptx.dml.color import RGBColor
from pptx.oxml.theme import CT_OfficeStyleSheet
from pptx.theme import ColorScheme, FontScheme, Theme

from .unitutil.mock import instance_mock


class DescribeTheme(object):
    """Unit-test suite for `pptx.theme.Theme` objects."""

    def it_provides_access_to_its_color_scheme(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        theme = Theme(theme_element, None)
        color_scheme = theme.color_scheme
        assert isinstance(color_scheme, ColorScheme)

    def it_provides_access_to_its_font_scheme(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        theme = Theme(theme_element, None)
        font_scheme = theme.font_scheme
        assert isinstance(font_scheme, FontScheme)

    def it_has_a_name_property(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        theme = Theme(theme_element, None)
        assert theme.name == "Office Theme"


class DescribeColorScheme(object):
    """Unit-test suite for `pptx.theme.ColorScheme` objects."""

    def it_can_read_accent1_as_RGBColor(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        cs = ColorScheme(theme_element.themeElements.clrScheme)
        assert cs.accent1 == RGBColor(0x4F, 0x81, 0xBD)

    @pytest.mark.parametrize(
        ("prop_name", "expected_hex"),
        [
            ("dark1", "000000"),
            ("light1", "FFFFFF"),
            ("dark2", "1F497D"),
            ("light2", "EEECE1"),
            ("accent1", "4F81BD"),
            ("accent2", "C0504D"),
            ("accent3", "9BBB59"),
            ("accent4", "8064A2"),
            ("accent5", "4BACC6"),
            ("accent6", "F79646"),
            ("hyperlink", "0000FF"),
            ("followed_hyperlink", "800080"),
        ],
    )
    def it_can_read_all_12_named_colors(self, prop_name, expected_hex):
        theme_element = CT_OfficeStyleSheet.new_default()
        cs = ColorScheme(theme_element.themeElements.clrScheme)
        color = getattr(cs, prop_name)
        assert color == RGBColor.from_string(expected_hex)

    def it_can_write_accent1(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        cs = ColorScheme(theme_element.themeElements.clrScheme)
        cs.accent1 = RGBColor(0xFF, 0x00, 0x00)
        assert cs.accent1 == RGBColor(0xFF, 0x00, 0x00)

    def it_handles_sysClr_elements(self):
        """dk1/lt1 in default theme use sysClr with lastClr fallback."""
        theme_element = CT_OfficeStyleSheet.new_default()
        cs = ColorScheme(theme_element.themeElements.clrScheme)
        # Read sysClr-based colors
        assert cs.dark1 == RGBColor(0x00, 0x00, 0x00)
        assert cs.light1 == RGBColor(0xFF, 0xFF, 0xFF)
        # Write replaces sysClr with srgbClr
        cs.dark1 = RGBColor(0x11, 0x22, 0x33)
        assert cs.dark1 == RGBColor(0x11, 0x22, 0x33)

    def it_has_a_name_property(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        cs = ColorScheme(theme_element.themeElements.clrScheme)
        assert cs.name == "Office"

    def it_can_iterate_colors(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        cs = ColorScheme(theme_element.themeElements.clrScheme)
        items = list(cs)
        assert len(items) == 12
        names = [name for name, _ in items]
        assert "dark1" in names
        assert "light1" in names
        assert "accent1" in names
        assert "hyperlink" in names
        assert "followed_hyperlink" in names
        # Verify items() gives same result
        assert cs.items() == items


class DescribeFontScheme(object):
    """Unit-test suite for `pptx.theme.FontScheme` objects."""

    def it_provides_major_font_name(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        fs = FontScheme(theme_element.themeElements.fontScheme)
        assert fs.major_font == "Calibri"

    def it_provides_minor_font_name(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        fs = FontScheme(theme_element.themeElements.fontScheme)
        assert fs.minor_font == "Calibri"

    def it_can_write_major_font(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        fs = FontScheme(theme_element.themeElements.fontScheme)
        fs.major_font = "Arial"
        assert fs.major_font == "Arial"

    def it_can_write_minor_font(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        fs = FontScheme(theme_element.themeElements.fontScheme)
        fs.minor_font = "Helvetica"
        assert fs.minor_font == "Helvetica"

    def it_has_a_name_property(self):
        theme_element = CT_OfficeStyleSheet.new_default()
        fs = FontScheme(theme_element.themeElements.fontScheme)
        assert fs.name == "Office"


class DescribeRoundTrip(object):
    """Integration test verifying theme modifications survive save/load cycle."""

    def it_round_trips_color_scheme_changes(self):
        from pptx import Presentation

        prs = Presentation()
        slide_master = prs.slide_masters[0]
        theme = slide_master.theme

        # Read original accent1
        original_accent1 = theme.color_scheme.accent1
        assert original_accent1 is not None

        # Modify accent1 to red
        theme.color_scheme.accent1 = RGBColor(0xFF, 0x00, 0x00)
        assert theme.color_scheme.accent1 == RGBColor(0xFF, 0x00, 0x00)

        # Save to BytesIO
        stream = BytesIO()
        prs.save(stream)

        # Reload and verify
        stream.seek(0)
        prs2 = Presentation(stream)
        theme2 = prs2.slide_masters[0].theme
        assert theme2.color_scheme.accent1 == RGBColor(0xFF, 0x00, 0x00)

    def it_round_trips_font_scheme_changes(self):
        from pptx import Presentation

        prs = Presentation()
        slide_master = prs.slide_masters[0]
        theme = slide_master.theme

        # Modify fonts
        theme.font_scheme.major_font = "Arial"
        theme.font_scheme.minor_font = "Helvetica"

        # Save to BytesIO
        stream = BytesIO()
        prs.save(stream)

        # Reload and verify
        stream.seek(0)
        prs2 = Presentation(stream)
        theme2 = prs2.slide_masters[0].theme
        assert theme2.font_scheme.major_font == "Arial"
        assert theme2.font_scheme.minor_font == "Helvetica"
